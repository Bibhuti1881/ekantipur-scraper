"""
=============================================================
  Ekantipur.com — News Scraper (Professional Edition)
  Author  : Bibhuti Adhikari
  Output  : ekantipur_news.xlsx + output.json
  Tools   : Playwright + OpenPyXL
=============================================================

Install:
    pip install playwright openpyxl
    playwright install chromium

Run:
    python ekantipur_scraper.py
"""

from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import json
import time

# ── Config ─────────────────────────────────────────────────────────────────────
OUTPUT_EXCEL = "ekantipur_news.xlsx"
OUTPUT_JSON  = "output.json"

SECTIONS = [
    {"name": "मनोरञ्जन",   "label": "Entertainment", "url": None,   "selector": "text=मनोरञ्जन"},
    {"name": "खेलकुद",     "label": "Sports",         "url": None,   "selector": "text=खेलकुद"},
    {"name": "राजनीति",    "label": "Politics",       "url": None,   "selector": "text=राजनीति"},
    {"name": "अर्थ",       "label": "Business",       "url": None,   "selector": "text=अर्थ"},
    {"name": "प्रविधि",    "label": "Technology",     "url": None,   "selector": "text=प्रविधि"},
]
ARTICLES_PER_SECTION = 10   # scrape 10 articles per section = 50 total


# ── Scrape one section ─────────────────────────────────────────────────────────
def scrape_section(page, section):
    articles_data = []
    label = section["label"]
    print(f"\n  📰 Scraping {label} ({section['name']}) ...")

    try:
        # Go to homepage and click section
        page.goto("https://ekantipur.com", timeout=20000)
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(1500)

        page.click(section["selector"])
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)

        # Try multiple article selectors
        selectors = [".category", "article", ".news-item", ".story-item", "[class*='article']"]
        articles = None
        for sel in selectors:
            try:
                page.wait_for_selector(sel, timeout=5000)
                articles = page.locator(sel)
                if articles.count() > 0:
                    break
            except:
                continue

        if not articles or articles.count() == 0:
            print(f"  ⚠️  No articles found for {label}")
            return articles_data

        count = articles.count()
        scraped = 0

        for i in range(count):
            if scraped >= ARTICLES_PER_SECTION:
                break

            article = articles.nth(i)

            try:
                item = {
                    "section_np":  section["name"],
                    "section_en":  label,
                    "title":       "N/A",
                    "author":      "N/A",
                    "date":        "N/A",
                    "image_url":   "N/A",
                    "article_url": "N/A",
                    "summary":     "N/A",
                    "scraped_at":  datetime.now().strftime("%Y-%m-%d %H:%M"),
                }

                # Title
                for t_sel in ["h2 a", "h3 a", "h1 a", ".title a", "a[href*='/']"]:
                    el = article.locator(t_sel)
                    if el.count() > 0:
                        item["title"] = el.first.text_content().strip()
                        # Article URL
                        href = el.first.get_attribute("href")
                        if href:
                            item["article_url"] = (
                                href if href.startswith("http")
                                else "https://ekantipur.com" + href
                            )
                        break

                if item["title"] in ("N/A", "", None):
                    continue

                # Author
                for a_sel in [".author-name a", ".author a", "[class*='author']", ".reporter"]:
                    el = article.locator(a_sel)
                    if el.count() > 0:
                        item["author"] = el.first.text_content().strip()
                        break

                # Date
                for d_sel in ["time", ".date", "[class*='date']", "[class*='time']"]:
                    el = article.locator(d_sel)
                    if el.count() > 0:
                        item["date"] = el.first.text_content().strip()
                        break

                # Image
                img = article.locator("img")
                if img.count() > 0:
                    img_url = (
                        img.first.get_attribute("src") or
                        img.first.get_attribute("data-src") or
                        img.first.get_attribute("srcset") or "N/A"
                    )
                    item["image_url"] = img_url

                # Summary / description
                for s_sel in ["p", ".summary", ".excerpt", "[class*='desc']"]:
                    el = article.locator(s_sel)
                    if el.count() > 0:
                        text = el.first.text_content().strip()
                        if len(text) > 10:
                            item["summary"] = text[:200]
                            break

                articles_data.append(item)
                scraped += 1

            except Exception as e:
                continue

        print(f"  ✅ {scraped} articles scraped from {label}")

    except Exception as e:
        print(f"  ❌ Error in {label}: {e}")

    return articles_data


# ── Scrape cartoon ─────────────────────────────────────────────────────────────
def scrape_cartoon(page):
    print("\n  🎨 Scraping Cartoon of the Day ...")
    cartoon = {
        "title":     "N/A",
        "image_url": "N/A",
        "author":    "N/A",
        "date":      datetime.now().strftime("%Y-%m-%d"),
    }
    try:
        page.goto("https://ekantipur.com/cartoon", timeout=15000)
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)

        # Image
        for sel in [".cartoon-image img", ".cartoon img", "article img", ".main-image img"]:
            el = page.locator(sel)
            if el.count() > 0:
                cartoon["image_url"] = (
                    el.first.get_attribute("src") or
                    el.first.get_attribute("data-src") or "N/A"
                )
                break

        # Title & author
        for sel in [".cartoon-description p", ".caption", ".description p", "figcaption"]:
            el = page.locator(sel)
            if el.count() > 0:
                text = el.first.text_content().strip()
                if "-" in text:
                    parts = text.split("-", 1)
                    cartoon["title"]  = parts[0].strip()
                    cartoon["author"] = parts[1].strip()
                else:
                    cartoon["title"] = text
                break

        print(f"  ✅ Cartoon scraped: {cartoon['title']}")
    except Exception as e:
        print(f"  ❌ Cartoon error: {e}")

    return cartoon


# ── Export to Excel ────────────────────────────────────────────────────────────
def export_to_excel(all_articles, cartoon, breaking_news=[]):
    print("\n📊 Creating Excel file ...")

    wb = Workbook()

    # ── Styles ─────────────────────────────────────────────────────────────────
    red_fill    = PatternFill("solid", fgColor="CC0000")   # ekantipur red
    dark_fill   = PatternFill("solid", fgColor="1A1A2E")
    alt_fill    = PatternFill("solid", fgColor="FFF5F5")   # light pink
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    head_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title_font  = Font(name="Calibri", bold=True, color="1A1A2E", size=10)
    data_font   = Font(name="Calibri", size=10)
    url_font    = Font(name="Calibri", size=9, color="0563C1", underline="single")
    center      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left        = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin        = Side(style="thin", color="DDDDDD")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 1 — All Articles
    # ══════════════════════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = "All Articles"

    # Title banner
    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value     = "📰  EKANTIPUR.COM — News Scraper (Professional Edition)"
    t.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
    t.fill      = red_fill
    t.alignment = center
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:J2")
    s = ws["A2"]
    s.value     = (f"Total articles: {len(all_articles)}   |   "
                   f"Sections: {len(SECTIONS)}   |   "
                   f"Source: ekantipur.com   |   "
                   f"Scraped: {datetime.now().strftime('%Y-%m-%d %H:%M')}   |   "
                   f"By: Bibhuti Adhikari")
    s.font      = Font(name="Calibri", italic=True, color="FFFFFF", size=10)
    s.fill      = dark_fill
    s.alignment = center
    ws.row_dimensions[2].height = 18

    # Headers
    headers    = ["#", "Section (NP)", "Section (EN)", "Title",
                  "Author", "Date", "Summary", "Image URL", "Article URL", "Scraped At"]
    col_widths = [4,   14,             14,              45,
                  18,       14,     35,          40,           40,             18]

    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = head_font; c.fill = dark_fill
        c.alignment = center; c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 22

    # Data
    for idx, item in enumerate(all_articles, start=1):
        row  = idx + 3
        fill = alt_fill if idx % 2 == 0 else white_fill
        data = [
            idx,
            item.get("section_np",  "N/A"),
            item.get("section_en",  "N/A"),
            item.get("title",       "N/A"),
            item.get("author",      "N/A"),
            item.get("date",        "N/A"),
            item.get("summary",     "N/A"),
            item.get("image_url",   "N/A"),
            item.get("article_url", "N/A"),
            item.get("scraped_at",  "N/A"),
        ]
        for ci, val in enumerate(data, start=1):
            c = ws.cell(row=row, column=ci, value=val)
            c.fill   = fill
            c.border = border
            if ci == 4:
                c.font = title_font; c.alignment = left
            elif ci in (8, 9):
                c.font = url_font; c.alignment = left
            else:
                c.font = data_font; c.alignment = center if ci in (1,2,3,5,6,10) else left
        ws.row_dimensions[row].height = 22

    ws.freeze_panes = "A4"

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 2 — By Section
    # ══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("By Section")

    ws2.merge_cells("A1:E1")
    h = ws2["A1"]
    h.value = "📂  ARTICLES BY SECTION"
    h.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    h.fill = red_fill; h.alignment = center
    ws2.row_dimensions[1].height = 28

    # Section headers
    for i, (h_text, w) in enumerate(zip(
        ["Section", "Nepali Name", "Articles Scraped", "Authors Found", "Has Images"],
        [20, 16, 18, 16, 14]
    ), start=1):
        c = ws2.cell(row=2, column=i, value=h_text)
        c.font = head_font; c.fill = dark_fill
        c.alignment = center; c.border = border
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[2].height = 20

    for ri, section in enumerate(SECTIONS, start=3):
        sec_articles = [a for a in all_articles if a["section_en"] == section["label"]]
        authors  = sum(1 for a in sec_articles if a["author"] != "N/A")
        images   = sum(1 for a in sec_articles if a["image_url"] != "N/A")
        row_fill = alt_fill if ri % 2 == 0 else white_fill

        for ci, val in enumerate([
            section["label"], section["name"],
            len(sec_articles), authors,
            f"{images}/{len(sec_articles)}"
        ], start=1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.font = data_font; c.fill = row_fill
            c.alignment = center; c.border = border
        ws2.row_dimensions[ri].height = 20

    ws2.freeze_panes = "A3"

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 3 — Cartoon of the Day
    # ══════════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("Cartoon of the Day")

    ws3.merge_cells("A1:B1")
    h3 = ws3["A1"]
    h3.value = "🎨  CARTOON OF THE DAY"
    h3.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    h3.fill = red_fill; h3.alignment = center
    ws3.row_dimensions[1].height = 28

    cartoon_rows = [
        ("Title",     cartoon.get("title",     "N/A")),
        ("Author",    cartoon.get("author",    "N/A")),
        ("Date",      cartoon.get("date",      "N/A")),
        ("Image URL", cartoon.get("image_url", "N/A")),
    ]
    for ri, (label, val) in enumerate(cartoon_rows, start=2):
        lc = ws3.cell(row=ri, column=1, value=label)
        vc = ws3.cell(row=ri, column=2, value=val)
        row_fill = alt_fill if ri % 2 == 0 else white_fill
        lc.font = Font(name="Calibri", bold=True, size=11)
        vc.font = url_font if label == "Image URL" else data_font
        lc.fill = row_fill
        vc.fill = PatternFill("solid", fgColor="FFF5F5") if ri % 2 == 0 else white_fill
        lc.alignment = left; vc.alignment = left
        lc.border = border; vc.border = border
        ws3.row_dimensions[ri].height = 20

    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 60

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 4 — Summary
    # ══════════════════════════════════════════════════════════════════════════
    ws4 = wb.create_sheet("Summary")

    ws4.merge_cells("A1:B1")
    hs = ws4["A1"]
    hs.value = "📊  SCRAPE SUMMARY"
    hs.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    hs.fill = red_fill; hs.alignment = center
    ws4.row_dimensions[1].height = 28

    authors_found = sum(1 for a in all_articles if a["author"] != "N/A")
    images_found  = sum(1 for a in all_articles if a["image_url"] != "N/A")
    urls_found    = sum(1 for a in all_articles if a["article_url"] != "N/A")

    summary = [
        ("Total Articles Scraped",  len(all_articles)),
        ("Total Sections",          len(SECTIONS)),
        ("Articles with Author",    authors_found),
        ("Articles with Image",     images_found),
        ("Articles with URL",       urls_found),
        ("Cartoon Scraped",         "Yes" if cartoon["title"] != "N/A" else "No"),
        ("Source",                  "ekantipur.com"),
        ("Scraped On",              datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Scraped By",              "Bibhuti Adhikari"),
        ("Tools Used",              "Python, Playwright, OpenPyXL"),
    ]

    for ri, (label, val) in enumerate(summary, start=2):
        lc = ws4.cell(row=ri, column=1, value=label)
        vc = ws4.cell(row=ri, column=2, value=val)
        row_fill = alt_fill if ri % 2 == 0 else white_fill
        lc.font = Font(name="Calibri", bold=True, size=11)
        vc.font = Font(name="Calibri", size=11)
        lc.fill = row_fill
        vc.fill = PatternFill("solid", fgColor="FFF5F5") if ri % 2 == 0 else white_fill
        lc.alignment = left; vc.alignment = center
        lc.border = border; vc.border = border
        ws4.row_dimensions[ri].height = 20

    ws4.column_dimensions["A"].width = 28
    ws4.column_dimensions["B"].width = 30
    ws4.freeze_panes = "A2"

    wb.save(OUTPUT_EXCEL)
    print(f"✅ Excel saved → {OUTPUT_EXCEL}")
    print(f"   • Sheet 1: All {len(all_articles)} articles")
    print(f"   • Sheet 2: By section breakdown")
    print(f"   • Sheet 3: Cartoon of the day")
    print(f"   • Sheet 4: Summary statistics")


# ── Main ───────────────────────────────────────────────────────────────────────
def scrape():
    all_articles = []

    print("\n📰 Ekantipur.com News Scraper — Starting ...")
    print("─" * 55)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1280, "height": 800}
        )
        page = context.new_page()

        # Scrape all sections
        for section in SECTIONS:
            articles = scrape_section(page, section)
            all_articles.extend(articles)
            time.sleep(1)

        # Scrape cartoon
        cartoon = scrape_cartoon(page)

        browser.close()

    print(f"\n✅ Total articles scraped: {len(all_articles)}")

    # Save JSON
    output = {
        "metadata": {
            "total_articles": len(all_articles),
            "sections":       len(SECTIONS),
            "scraped_at":     datetime.now().strftime("%Y-%m-%d %H:%M"),
            "scraped_by":     "Bibhuti Adhikari",
            "source":         "ekantipur.com",
        },
        "articles":          all_articles,
        "cartoon_of_the_day": cartoon,
    }

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)
    print(f"✅ JSON saved → {OUTPUT_JSON}")

    # Save Excel
    if all_articles:
        export_to_excel(all_articles, cartoon)

    print("\n🎉 Done! Open ekantipur_news.xlsx to see your data.")


if __name__ == "__main__":
    scrape()